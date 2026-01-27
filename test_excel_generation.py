"""
Simple test script to verify Excel generation works
"""
from openpyxl import Workbook
import io

# Create a simple workbook
wb = Workbook()
ws = wb.active
ws.title = "Test"

# Add some data
ws['A1'] = 'Nombre'
ws['B1'] = 'Email'
ws['A2'] = 'Juan Pérez'
ws['B2'] = 'juan@example.com'

# Save to file
output = io.BytesIO()
wb.save(output)
output.seek(0)

# Save to disk for testing
with open('test_export.xlsx', 'wb') as f:
    f.write(output.read())

print("✅ Excel file created successfully: test_export.xlsx")
print(f"File size: {len(output.getvalue())} bytes")
